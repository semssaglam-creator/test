#!/usr/bin/env python3
"""Takip listesi PDF -> Excel dönüştürücü.

Sütun çizgisi olmayan, metni sütun konumlarına göre yazılmış takip listesi
PDF'lerini okur; her kaydı 15 sütuna ayırıp .xlsx olarak kaydeder.

Kullanım:
    python3 takip2excel.py                 # pencereli arayüz
    python3 takip2excel.py a.pdf [b.pdf]   # komut satırı (-o cikti.xlsx)
"""
import argparse
import re
import sys
from collections import Counter
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RIGHT_COLS = [
    "Ana Vergi Kodu", "Vergi Kodu", "Ana Takip Dosya No", "Takip Dosya No",
    "Vergi Aslı Borcu Toplamı", "KGZ Toplamı", "Ceza Tutarı", "Toplam Borç",
]
COLUMNS = ["Sıra No", "Vergi No", "TC Kimlik No", "Plaka No", "Soyad Ad/Unvanı",
           "Adres", "Vergi Dönemi", *RIGHT_COLS]
AMOUNT_COLS = RIGHT_COLS[4:]
TEXT_COLS = {"Vergi No", "TC Kimlik No", "Ana Vergi Kodu", "Vergi Kodu",
             "Ana Takip Dosya No", "Takip Dosya No"}

PERIOD_IN = re.compile(r"\d{2}/\d{4}-\d{2}/\d{4}")
PERIOD = re.compile(rf"^{PERIOD_IN.pattern}$")
INT = re.compile(r"^\d{1,6}$")
VKN = re.compile(r"^\d{10}$")
TCKN = re.compile(r"^\d{11}$")
PLAKA = re.compile(r"^\d{2}[A-ZÇĞİÖŞÜ]{1,3}\d{2,5}$")
AMOUNT = re.compile(r"^-?\d{1,3}(\.\d{3})*,\d{2}$")
HEADER_WORDS = {
    "sıra", "no", "vergi", "tc", "kimlik", "plaka", "soyad", "ad/unvanı",
    "adres", "dönemi", "ana", "kodu", "(*)", "takip", "dosya", "aslı", "borcu",
    "toplamı", "kgz", "ceza", "tutarı", "toplam", "borç",
}


def _is_header_line(words):
    hits = sum(w["text"].lower().replace("i̇", "i") in HEADER_WORDS for w in words)
    return hits >= max(2, len(words) * 0.6)


def _lines(words, tol=2.0):
    """Kelimeleri satırlara grupla (üstten alta, soldan sağa)."""
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(w["top"] - lines[-1][0]["top"]) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(l, key=lambda w: w["x0"]) for l in lines]


def _join(words):
    return " ".join(" ".join(w["text"] for w in l) for l in _lines(words))


def _cut_at_gap(words, max_gap):
    """Satırlar arası boşluk max_gap'i aşınca kes (alt bilgi vb. dışarıda kalsın)."""
    out, prev = [], None
    for line in _lines(words):
        if prev is not None and line[0]["top"] - prev > max_gap:
            break
        out.extend(line)
        prev = line[0]["top"]
    return out


def _mode(values, step=2.0):
    if not values:
        return None
    bucket, _ = Counter(round(v / step) for v in values).most_common(1)[0]
    return min(v for v in values if round(v / step) == bucket)


def _split_fused(words):
    """Bitişik sütunlar tek kelime gelirse (ör. 'OĞLU11/2021-11/2021') dönemden ayır."""
    out = []
    for w in words:
        m = PERIOD_IN.search(w["text"])
        if not m or (m.start() == 0 and m.end() == len(w["text"])):
            out.append(w)
            continue
        n, width = len(w["text"]), w["x1"] - w["x0"]
        for s, e in ((0, m.start()), (m.start(), m.end()), (m.end(), n)):
            if s < e:
                out.append({**w, "text": w["text"][s:e],
                            "x0": w["x0"] + width * s / n, "x1": w["x0"] + width * e / n})
    return out


def _collect_records(pdf):
    """Her kayıt için ham kelime kümesini döndür. Çapa: Vergi Dönemi hücresi."""
    records = []
    for page in pdf.pages:
        words = _split_fused(page.extract_words(x_tolerance=1.5, y_tolerance=2))
        anchors = sorted((w for w in words if PERIOD.match(w["text"])), key=lambda w: w["top"])
        if not anchors:
            continue
        h = anchors[0]["bottom"] - anchors[0]["top"]
        eps = h * 0.5

        # Sayfa başında, başlık bloğunun altında kalan satırlar önceki kaydın devamıdır.
        pre = [w for w in words if w["top"] < anchors[0]["top"] - eps]
        header_bottom = max((l[-1]["bottom"] for l in _lines(pre) if _is_header_line(l)), default=None)
        if records and header_bottom is not None:
            cont = [w for w in pre if w["top"] > header_bottom]
            records[-1]["cont"].extend(_cut_at_gap(cont, h * 2.5))

        for i, a in enumerate(anchors):
            top = a["top"] - eps
            bottom = anchors[i + 1]["top"] - eps if i + 1 < len(anchors) else page.height
            body = [w for w in words if top <= w["top"] < bottom]
            records.append({"anchor": a, "h": h, "words": _cut_at_gap(body, h * 2.5),
                            "cont": [], "page": page.page_number})
    return records


def parse_pdf(path):
    """PDF'i oku, (satırlar, uyarılar) döndür. Satır = COLUMNS sırasında liste."""
    with pdfplumber.open(path) as pdf:
        records = _collect_records(pdf)
    if not records:
        raise ValueError("PDF'te kayıt bulunamadı (metin katmanı yok veya biçim farklı).")

    # --- Kalibrasyon: ad sütunu sol kenarı (N0) ve adres sütunu sol kenarı (A0)
    name_starts = []
    for r in records:
        a = r["anchor"]
        first = [w for w in r["words"] if abs(w["top"] - a["top"]) <= r["h"] * 0.6 and w["x1"] < a["x0"]]
        text = [w for w in sorted(first, key=lambda w: w["x0"])
                if not (INT.match(w["text"]) or VKN.match(w["text"]) or TCKN.match(w["text"])
                        or PLAKA.match(w["text"]))]
        if text:
            name_starts.append(text[0]["x0"])
        r["first"] = first
    n0 = _mode(name_starts)
    if n0 is None:
        raise ValueError("Ad/Unvan sütunu tespit edilemedi.")

    zone = [w["x0"] for r in records for w in r["words"] + r["cont"]
            if n0 + 15 < w["x0"] and w["x1"] <= r["anchor"]["x0"]]
    a0 = _mode(zone)

    # --- Kalibrasyon: dönemin sağındaki 8 sütunun merkezleri
    def right_tokens(r):
        return sorted((w for w in r["words"] if w["x0"] > r["anchor"]["x1"]), key=lambda w: w["x0"])

    full = [right_tokens(r) for r in records]
    full = [t for t in full if len(t) == len(RIGHT_COLS)]
    centers = [sum((t[i]["x0"] + t[i]["x1"]) / 2 for t in full) / len(full)
               for i in range(len(RIGHT_COLS))] if full else None

    rows, warnings = [], []
    for r in records:
        a = r["anchor"]
        row = dict.fromkeys(COLUMNS, "")
        row["Vergi Dönemi"] = a["text"]

        plaka = []
        for w in r["first"]:
            t = w["text"]
            if w["x0"] >= n0 - 2:
                continue
            if not row["Sıra No"] and INT.match(t):
                row["Sıra No"] = int(t)
            elif VKN.match(t):
                row["Vergi No"] = t
            elif TCKN.match(t):
                row["TC Kimlik No"] = t
            else:
                plaka.append(w)
        row["Plaka No"] = _join(plaka)

        text = [w for w in r["words"] + r["cont"] if n0 - 2 <= w["x0"] and w["x1"] <= a["x0"] + 1]
        split = a0 - 2 if a0 is not None else float("inf")
        row["Soyad Ad/Unvanı"] = _join([w for w in text if w["x0"] < split])
        row["Adres"] = _join([w for w in text if w["x0"] >= split])

        toks = right_tokens(r)
        cells = {c: [] for c in RIGHT_COLS}
        if len(toks) == len(RIGHT_COLS) or not centers:
            for c, w in zip(RIGHT_COLS, toks):
                cells[c].append(w)
        else:
            for w in toks:
                cx = (w["x0"] + w["x1"]) / 2
                i = min(range(len(centers)), key=lambda i: abs(centers[i] - cx))
                cells[RIGHT_COLS[i]].append(w)
        for c in RIGHT_COLS:
            row[c] = " ".join(w["text"] for w in cells[c])

        ref = f"sayfa {r['page']}, sıra {row['Sıra No'] or '?'}"
        for c in AMOUNT_COLS:
            v = row[c]
            if AMOUNT.match(v):
                row[c] = float(v.replace(".", "").replace(",", "."))
            elif v:
                warnings.append(f"{ref}: '{c}' sayı değil: {v!r}")
        if all(isinstance(row[c], float) for c in AMOUNT_COLS):
            asli, kgz, ceza, top = (row[c] for c in AMOUNT_COLS)
            if abs(asli + kgz + ceza - top) > 0.005:
                warnings.append(f"{ref}: Toplam Borç ({top:.2f}) ≠ Aslı+KGZ+Ceza ({asli + kgz + ceza:.2f})")
        if len(toks) != len(RIGHT_COLS):
            warnings.append(f"{ref}: sağ tarafta {len(toks)} değer var (beklenen {len(RIGHT_COLS)}), kontrol edin")
        if not row["Vergi No"] and not row["TC Kimlik No"]:
            warnings.append(f"{ref}: Vergi No / TC Kimlik No bulunamadı")
        rows.append([row[c] for c in COLUMNS])
    return rows, warnings


def write_xlsx(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Takip Listesi"
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    amount_idx = {COLUMNS.index(c) + 1 for c in AMOUNT_COLS}
    text_idx = {COLUMNS.index(c) + 1 for c in TEXT_COLS}
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in r:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (5, 6))
            if cell.column in amount_idx:
                cell.number_format = "#,##0.00"
            elif cell.column in text_idx:
                cell.number_format = "@"

    last = ws.max_row
    total = last + 1
    ws.cell(total, COLUMNS.index("Vergi Dönemi") + 1, "TOPLAM").font = Font(bold=True)
    for c in amount_idx:
        col = get_column_letter(c)
        cell = ws.cell(total, c, f"=SUM({col}2:{col}{last})")
        cell.number_format = "#,##0.00"
        cell.font = Font(bold=True)

    widths = [7, 12, 13, 10, 32, 45, 16, 9, 9, 22, 22, 14, 11, 11, 13]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    wb.save(out_path)


def convert(pdf_paths, out_path, log=print):
    all_rows = []
    for p in pdf_paths:
        rows, warnings = parse_pdf(p)
        log(f"{Path(p).name}: {len(rows)} kayıt okundu")
        for w in warnings:
            log(f"  UYARI {w}")
        all_rows.extend(rows)
    write_xlsx(all_rows, out_path)
    log(f"Kaydedildi: {out_path} ({len(all_rows)} kayıt)")
    return len(all_rows)


def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("Takip PDF → Excel")
    root.geometry("760x480")
    files = []

    top = tk.Frame(root, padx=10, pady=10)
    top.pack(fill="x")
    label = tk.Label(top, text="PDF seçilmedi", anchor="w")

    def log(msg):
        box.insert("end", msg + "\n")
        box.see("end")
        root.update_idletasks()

    def pick():
        sel = filedialog.askopenfilenames(title="PDF seç", filetypes=[("PDF", "*.pdf *.PDF")])
        if sel:
            files[:] = sel
            label.config(text=f"{len(sel)} dosya: " + ", ".join(Path(s).name for s in sel))

    def run():
        if not files:
            messagebox.showwarning("Uyarı", "Önce PDF seçin.")
            return
        first = Path(files[0])
        out = filedialog.asksaveasfilename(
            title="Excel kaydet", defaultextension=".xlsx", initialdir=first.parent,
            initialfile=first.stem + ".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not out:
            return
        box.delete("1.0", "end")
        try:
            n = convert(files, out, log)
            messagebox.showinfo("Tamam", f"{n} kayıt Excel'e aktarıldı.\n{out}")
        except Exception as e:
            log(f"HATA: {e}")
            messagebox.showerror("Hata", str(e))

    tk.Button(top, text="PDF Seç…", command=pick, width=14).pack(side="left")
    tk.Button(top, text="Excel'e Aktar", command=run, width=14).pack(side="left", padx=8)
    label.pack(side="left", fill="x", expand=True)
    box = scrolledtext.ScrolledText(root, font=("monospace", 9))
    box.pack(fill="both", expand=True, padx=10, pady=(0, 10))
    root.mainloop()


def main():
    ap = argparse.ArgumentParser(description="Takip listesi PDF'ini Excel'e aktarır.")
    ap.add_argument("pdf", nargs="*", help="PDF dosya(lar)ı; verilmezse pencere açılır")
    ap.add_argument("-o", "--output", help="Çıktı .xlsx (varsayılan: ilk PDF adıyla)")
    args = ap.parse_args()
    if not args.pdf:
        run_gui()
        return
    out = args.output or str(Path(args.pdf[0]).with_suffix(".xlsx"))
    try:
        convert(args.pdf, out)
    except Exception as e:
        sys.exit(f"HATA: {e}")


if __name__ == "__main__":
    main()
