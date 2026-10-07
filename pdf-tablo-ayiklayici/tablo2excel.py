#!/usr/bin/env python3
"""
tablo2excel.py — Çizgisiz PDF tablosunu (amme alacakları borç listesi) Excel'e çevirir.

Kullanım:
    python3 tablo2excel.py liste.pdf                 # liste.xlsx üretir
    python3 tablo2excel.py liste.pdf -o sonuc.xlsx
    python3 tablo2excel.py liste.pdf --sayfalar 1-3  # yalnızca bu sayfalar

Gerekenler (bir kez):  pip install pdfplumber openpyxl
Her şey bilgisayarında çalışır; dosya hiçbir yere gönderilmez.
"""
import argparse
import datetime as dt
import os
import re
import sys

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Eksik paket var. Önce şunu çalıştır:\n    pip install pdfplumber openpyxl\n"
             "(Sistem izin vermezse: pip install --user pdfplumber openpyxl)")

BASLIKLAR = ["Sıra No", "Vergi No", "TC Kimlik No", "Plaka No", "Adı Soyadı / Ünvanı", "Adres",
             "Vergi Dönemi", "Ana Vergi Kodu", "Vergi Kodu", "Ana Takip Dosya No", "Takip Dosya No",
             "Vergi Aslı Borcu Toplamı", "KGZ Toplamı", "Ceza Tutarı", "Toplam Borç"]
# Bir borcun alt kalemlerinde boş kalırsa üst satırdan doldurulacak sütunlar
DOLDUR = ["Sıra No", "Vergi No", "TC Kimlik No", "Adı Soyadı / Ünvanı", "Adres"]

SATIR_TOL = 2.0      # aynı satır sayılacak dikey sapma (pt)
KELIME_ARALIGI = 0.4  # bunun altındaki boşluk (yazı boyunun katı) aynı hücre sayılır
GURULTU = 0.03       # sütun koridorunu bu oranda satırın kesmesine izin ver

NUM_TR = re.compile(r"^[-−(]?\d{1,3}(\.\d{3})*(,\d+)?\)?$|^[-−(]?\d+,\d+\)?$")
TARIH = re.compile(r"^(\d{2})[./](\d{2})[./](\d{4})$")


# ---------- PDF'ten kelimeler ----------
def kelimeler(page):
    ws = page.extract_words(x_tolerance=1.5, y_tolerance=2, keep_blank_chars=False, use_text_flow=False,
                            extra_attrs=["size"])
    return [{"x": w["x0"], "w": w["x1"] - w["x0"], "y": w["bottom"], "fs": w.get("size") or 8, "s": w["text"]}
            for w in ws if w["text"].strip()]


def satirlar(items):
    items = sorted(items, key=lambda a: (a["y"], a["x"]))
    lines = []
    for it in items:
        if lines and abs(it["y"] - lines[-1]["y"]) <= SATIR_TOL:
            L = lines[-1]
            L["items"].append(it)
            L["y"] = sum(i["y"] for i in L["items"]) / len(L["items"])
        else:
            lines.append({"y": it["y"], "items": [it]})
    for L in lines:
        L["items"].sort(key=lambda a: a["x"])
        segs, cur = [], None
        for it in L["items"]:
            if cur and it["x"] - (cur["x"] + cur["w"]) < KELIME_ARALIGI * max(cur["fs"], it["fs"]):
                cur["w"] = max(cur["x"] + cur["w"], it["x"] + it["w"]) - cur["x"]
                cur["s"] += " " + it["s"]
                cur["fs"] = max(cur["fs"], it["fs"])
                continue
            cur = dict(it)
            segs.append(cur)
        L["segs"] = segs
    return lines


def paragraf_isaretle(lines, W):
    """Art arda en az 2 geniş düz metin satırı (tutanak, açıklama) paragraf sayılır."""
    wide = [len(L["segs"]) <= 2 and any(s["w"] > W * 0.45 for s in L["segs"]) for L in lines]
    for k, L in enumerate(lines):
        L["prose"] = wide[k] and ((k > 0 and wide[k - 1]) or (k + 1 < len(lines) and wide[k + 1]))
    for k, L in enumerate(lines):
        if not L["prose"] and k > 0 and lines[k - 1]["prose"] and wide[k - 1] and len(L["segs"]) <= 2 \
                and not any(re.search(r"\d,\d{2}\b", s["s"]) for s in L["segs"]):
            L["prose"] = True


# ---------- Sütun sınırları ----------
def col_of(x, divs):
    k = 0
    while k < len(divs) and x > divs[k]:
        k += 1
    return k


def sinirlar(pages_lines, W, hedef):
    W = int(W) + 2
    cov = [0] * W
    n, lo, hi = 0, W, 0
    xs = []
    for lines in pages_lines:
        for L in lines:
            if L["prose"] or (len(L["segs"]) <= 2 and any(s["w"] > W * 0.45 for s in L["segs"])):
                continue
            n += 1
            for it in L["items"]:
                a, b = max(0, round(it["x"])), min(W - 1, round(it["x"] + it["w"]) - 1)
                for i in range(a, b + 1):
                    cov[i] += 1
                lo, hi = min(lo, a), max(hi, b)
                xs.append(it["x"] + it["w"] / 2)
    if not n:
        return []
    allow = max(1, round(n * GURULTU))
    runs, start = [], -1
    for i in range(lo, hi + 2):
        empty = i <= hi and cov[i] <= allow
        if empty and start < 0:
            start = i
        if not empty and start >= 0:
            runs.append((start, i))
            start = -1
    runs.sort(key=lambda r: r[1] - r[0], reverse=True)
    divs = []
    for a, b in runs:
        if len(divs) >= hedef - 1:
            break
        trial = sorted(divs + [(a + b) / 2])
        cnt = [0] * (len(trial) + 1)
        for x in xs:
            cnt[col_of(x, trial)] += 1
        if all(c > allow for c in cnt):
            divs = trial
    return divs


# ---------- Satırları hücrelere böl ----------
def sayi(s):
    t = re.sub(r"\s", "", s)
    if not NUM_TR.match(t):
        if re.fullmatch(r"-?\d+", t) and not re.match(r"-?0\d", t) and len(t.lstrip("-")) < 10:
            return int(t)
        return None
    neg = bool(re.match(r"^[-−(]", t) or t.endswith(")"))
    core = re.sub(r"^[-−(]|\)$", "", t)
    if re.match(r"0\d", core) and "," not in core:
        return None
    v = float(core.replace(".", "").replace(",", "."))
    return -v if neg else v


def yapistir(a, b):
    """Alta kayan kod/dönem parçaları boşluksuz birleşir: 2026010166Evn + 0000001, 01/2025- + 12/2025"""
    if re.search(r"[-/]$", a):
        return True
    return " " not in a and " " not in b and len(a) >= 8 and sum(ch.isdigit() for ch in a) >= 4 \
        and re.fullmatch(r"[0-9A-Za-z]+", b) is not None


def hucreler(L, divs, nc):
    cells, last = [""] * nc, [None] * nc
    for it in L["items"]:
        k = col_of(it["x"] + it["w"] / 2, divs)
        p = last[k]
        if not cells[k]:
            cells[k] = it["s"]
        else:
            cells[k] += (" " if it["x"] - (p["x"] + p["w"]) > 0.12 * it["fs"] else "") + it["s"]
        last[k] = it
    return cells


def sayfa_satirlari(lines, divs, nc, anahtar, doldur):
    # tablo başlangıcı: anahtar sütunda yalnızca sayı olan ilk satır
    data = [L for L in lines if not L["prose"]]
    first = next((i for i, L in enumerate(data) if re.fullmatch(r"\d+", hucreler(L, divs, nc)[anahtar].replace(" ", ""))), None)
    if first is None:
        return []
    rows = []
    for L in data[first:]:
        r = hucreler(L, divs, nc)
        if not any(r):
            continue
        tutar_var = any(v and re.search(r",\d+\)?$", v) and sayi(v) is not None for v in r)
        if rows and not r[anahtar] and not tutar_var:
            prev = rows[-1]
            for i, v in enumerate(r):
                if v:
                    prev[i] = prev[i] + ("" if yapistir(prev[i], v) else " ") + v if prev[i] else v
        else:
            rows.append(r)
    for i in range(1, len(rows)):
        if not rows[i][anahtar]:
            for c in doldur:
                if c < nc and not rows[i][c]:
                    rows[i][c] = rows[i - 1][c]
    return rows


# ---------- Excel ----------
def sutun_tipleri(rows, nc):
    tip = []
    for c in range(nc):
        vals = [r[c] for r in rows if r[c]]
        if vals and all(TARIH.match(v) for v in vals):
            tip.append("d")
        elif vals and all(sayi(v) is not None for v in vals):
            tip.append("n")
        else:
            tip.append("s")
    return tip


def excel_yaz(rows, basliklar, cikti):
    nc = len(basliklar)
    tip = sutun_tipleri(rows, nc)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tablo"
    ws.append(basliklar)
    for c in range(nc):
        ws.cell(row=1, column=c + 1).font = Font(bold=True)
    genislik = [len(h) + 2 for h in basliklar]
    for ri, r in enumerate(rows, start=2):
        for c, v in enumerate(r):
            if not v:
                continue
            cell = ws.cell(row=ri, column=c + 1)
            if tip[c] == "d":
                g, a, y = map(int, TARIH.match(v).groups())
                try:
                    cell.value = dt.date(y, a, g)
                    cell.number_format = "dd.mm.yyyy"
                except ValueError:
                    cell.value = v
            elif tip[c] == "n":
                cell.value = sayi(v)
                cell.number_format = "#,##0.00" if "," in v else "0"
            else:
                cell.value = v
                cell.number_format = "@"
            genislik[c] = min(60, max(genislik[c], len(v) + 1))
    for c, w in enumerate(genislik):
        ws.column_dimensions[get_column_letter(c + 1)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(cikti)


def sayfa_araligi(text, n):
    if not text:
        return list(range(n))
    out = set()
    for p in re.split(r"[,;\s]+", text.strip()):
        m = re.fullmatch(r"(\d+)(-)?(\d+)?", p)
        if not m:
            continue
        a = int(m.group(1))
        b = int(m.group(3)) if m.group(3) else (n if m.group(2) else a)
        out.update(range(max(1, min(a, b)) - 1, min(n, max(a, b))))
    return sorted(out) or list(range(n))


def main():
    ap = argparse.ArgumentParser(description="Çizgisiz PDF tablosunu Excel'e çevirir.")
    ap.add_argument("pdf", help="PDF dosyası")
    ap.add_argument("-o", "--cikti", help="Excel dosyasının adı (varsayılan: PDF ile aynı ad, .xlsx)")
    ap.add_argument("--sayfalar", help="ör. 1-3,7 ya da 2-  (varsayılan: tümü)")
    ap.add_argument("--parola", help="PDF parolalıysa", default="")
    a = ap.parse_args()

    cikti = a.cikti or os.path.splitext(a.pdf)[0] + ".xlsx"
    nc = len(BASLIKLAR)
    anahtar = 0
    doldur = [BASLIKLAR.index(h) for h in DOLDUR]
    try:
        pdf = pdfplumber.open(a.pdf, password=a.parola or None)
    except Exception as e:
        sys.exit(f"PDF açılamadı: {e}\n(Parolalıysa --parola ile ver.)")
    with pdf:
        secili = sayfa_araligi(a.sayfalar, len(pdf.pages))
        sayfalar, W = [], 0
        for i in secili:
            pg = pdf.pages[i]
            print(f"\rSayfa {i + 1}/{len(pdf.pages)} okunuyor…", end="", file=sys.stderr, flush=True)
            lines = satirlar(kelimeler(pg))
            paragraf_isaretle(lines, pg.width)
            sayfalar.append(lines)
            W = max(W, pg.width)
        print(file=sys.stderr)
    if not any(sayfalar):
        sys.exit("PDF'te seçilebilir metin yok (taranmış görüntü olabilir).")
    divs = sinirlar(sayfalar, W, nc)
    if len(divs) != nc - 1:
        print(f"Uyarı: {nc} yerine {len(divs) + 1} sütun bulunabildi; sonuçları kontrol et.", file=sys.stderr)
    nc2 = len(divs) + 1
    rows = []
    for lines in sayfalar:
        rows += sayfa_satirlari(lines, divs, nc2, anahtar, doldur)
    basliklar = (BASLIKLAR + [f"Sütun {i + 1}" for i in range(nc2)])[:nc2]
    excel_yaz(rows, basliklar, cikti)
    kayit = len({r[0] for r in rows if r[0]})
    print(f"Tamam: {len(rows)} satır ({kayit} kayıt), {nc2} sütun → {cikti}")


if __name__ == "__main__":
    main()
