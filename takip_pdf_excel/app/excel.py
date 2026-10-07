"""Kayıtlardan .xlsx üretir (bellekte; geçici dosya yok)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .ayristirici import AMOUNT_COLS, COLUMNS

EK_SUTUNLAR = ["Kontrol", "Kaynak PDF"]
METIN_SUTUNLARI = {"Vergi No", "TC Kimlik No", "Ana Vergi Kodu", "Vergi Kodu",
                   "Ana Takip Dosya No", "Takip Dosya No"}
GENISLIK = [7, 12, 13, 10, 32, 45, 16, 9, 9, 22, 22, 14, 11, 11, 13, 40, 20]

KIRMIZI = PatternFill("solid", fgColor="F8CBAD")
SATIR_UYARI = PatternFill("solid", fgColor="FCE4D6")
BASLIK = PatternFill("solid", fgColor="305496")


def excel_olustur(kayitlar):
    """kayitlar: [{"degerler", "uyarilar", "sorunlu", "dosya"}] -> xlsx baytları."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Takip Listesi"
    basliklar = COLUMNS + EK_SUTUNLAR
    ws.append(basliklar)

    ince = Side(style="thin", color="BFBFBF")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = BASLIK
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = kenar

    tutar_idx = {COLUMNS.index(c) + 1 for c in AMOUNT_COLS}
    metin_idx = {COLUMNS.index(c) + 1 for c in METIN_SUTUNLARI}
    sarma_idx = {COLUMNS.index("Soyad Ad/Unvanı") + 1, COLUMNS.index("Adres") + 1, len(COLUMNS) + 1}

    for k in kayitlar:
        uyarilar = k.get("uyarilar") or []
        sorunlu = {COLUMNS.index(c) + 1 for c in (k.get("sorunlu") or []) if c in COLUMNS}
        ws.append(list(k["degerler"]) + ["; ".join(uyarilar), k.get("dosya", "")])
        r = ws.max_row
        for cell in ws[r]:
            cell.border = kenar
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in sarma_idx)
            if cell.column in tutar_idx:
                cell.number_format = "#,##0.00"
            elif cell.column in metin_idx:
                cell.number_format = "@"
            if cell.column in sorunlu:
                cell.fill = KIRMIZI
            elif uyarilar:
                cell.fill = SATIR_UYARI
        if uyarilar:
            ws.cell(r, len(COLUMNS) + 1).font = Font(color="C00000", bold=True)

    son = ws.max_row
    t = son + 1
    ws.cell(t, COLUMNS.index("Vergi Dönemi") + 1, "TOPLAM").font = Font(bold=True)
    for c in sorted(tutar_idx):
        harf = get_column_letter(c)
        cell = ws.cell(t, c, "=SUM({0}2:{0}{1})".format(harf, son))
        cell.number_format = "#,##0.00"
        cell.font = Font(bold=True)
    sorunlu_sayisi = sum(1 for k in kayitlar if k.get("uyarilar"))
    if sorunlu_sayisi:
        not_hucre = ws.cell(t, len(COLUMNS) + 1,
                            "{} kayıtta uyarı var; kırmızı hücreleri PDF ile karşılaştırın. "
                            "Boş tutarlar toplama girmez.".format(sorunlu_sayisi))
        not_hucre.font = Font(color="C00000", bold=True)
        not_hucre.alignment = Alignment(wrap_text=True, vertical="top")

    for i, g in enumerate(GENISLIK, 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(basliklar)), son)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    akis = io.BytesIO()
    wb.save(akis)
    return akis.getvalue()
